"""Real file exports: an Excel workbook via openpyxl and a PDF via Playwright."""
import base64, html, json
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from playwright.async_api import async_playwright
from . import db, vault
from .config import EXPORTS

HEAD = Font(bold=True, color="FFFFFF", size=10)
FILL = PatternFill("solid", fgColor="123A6B")
STAGE_NAMES = {"lead": "Lead captured", "docs": "Documents pending", "register": "Portal registration",
               "feasible": "DISCOM feasibility", "netmeter": "Net metering", "install": "Installation",
               "claim": "Subsidy claim filed", "disbursed": "Subsidy disbursed"}


def _sheet(ws, headers, rows, widths):
    ws.append(headers)
    for i, c in enumerate(ws[1], start=1):
        c.font = HEAD
        c.fill = FILL
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[c.column_letter].width = widths[i - 1]
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"


def pipeline_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "Pipeline"
    _sheet(ws, ["Ref", "Customer", "Location", "Consumer no.", "kW", "Category", "Stage",
                "Application no.", "Subsidy (INR)", "Updated"],
           [[c["id"], c["name"], c["place"], c["consumer_no"], c["kw"], c["category"],
             STAGE_NAMES.get(c["stage"], c["stage"]), c["application_no"] or "-", c["subsidy"],
             c["last_activity"]] for c in db.q("SELECT * FROM customers ORDER BY id")],
           [10, 30, 26, 16, 7, 16, 22, 22, 14, 11])

    ws2 = wb.create_sheet("Numbers")
    _sheet(ws2, ["Number", "Assigned to", "Expires", "OTPs received", "Status"],
           [[n["msisdn"], n["customer_id"] or "-", (n["expires_at"] or "")[:10], n["otp_count"],
             n["status"]] for n in db.q("SELECT * FROM numbers ORDER BY id")],
           [20, 14, 14, 15, 12])

    ws3 = wb.create_sheet("Runs")
    _sheet(ws3, ["Run", "Customer", "Workflow", "Status", "Started", "Finished", "Reference",
                 "Retries"],
           [[r["id"], r["customer_id"], r["workflow"], r["status"],
             (r["started_at"] or "").replace("T", " "),
             (r["finished_at"] or "").replace("T", " "), r["application_no"] or "-", r["retries"]]
            for r in db.q("SELECT * FROM runs ORDER BY id DESC")],
           [7, 14, 26, 11, 20, 20, 22, 9])

    ws4 = wb.create_sheet("Audit")
    _sheet(ws4, ["Timestamp", "Actor", "Role", "Action", "Target", "Detail"],
           [[a["ts"].replace("T", " "), a["actor"], a["role"], a["action"], a["target"], a["detail"]]
            for a in db.q("SELECT * FROM audit ORDER BY id DESC")],
           [20, 14, 12, 34, 14, 60])

    path = EXPORTS / f"solara_pipeline_{datetime.now():%Y%m%d_%H%M}.xlsx"
    wb.save(path)
    return path


def _run_report_html(run_id):
    r = db.one("SELECT * FROM runs WHERE id=?", (run_id,))
    c = db.one("SELECT * FROM customers WHERE id=?", (r["customer_id"],))
    steps = db.q("SELECT * FROM run_steps WHERE run_id=? ORDER BY idx", (run_id,))
    events = db.q("SELECT * FROM run_events WHERE run_id=? ORDER BY id", (run_id,))
    shots = db.q("SELECT * FROM screenshots WHERE run_id=? ORDER BY id", (run_id,))
    holds = db.q("SELECT * FROM holds WHERE run_id=? ORDER BY id", (run_id,))
    rows = "".join(
        f"<tr><td>{s['idx'] + 1}</td><td>{html.escape(s['label'])}</td>"
        f"<td class=s-{s['status']}>{s['status']}</td><td>{html.escape(s['note'] or '')}</td><td class=n>{s['ms'] / 1000:.1f}s</td></tr>" for s in steps)
    hrows = "".join(
        f"<tr><td>{h['kind']}</td><td>{h['status']}</td><td>{h['resolved_by'] or '-'}</td>"
        f"<td>{h['resolution'] or '-'}</td><td class=n>{(h['resolved_at'] or '')[11:]}</td></tr>"
        for h in holds)
    log = "".join(f"<div class=le><span>{e['ts']}</span>{html.escape(e['message'])}</div>"
                  for e in events)
    # screenshots are encrypted on disk, so they are decrypted and embedded inline
    imgs = "".join(
        f"<figure><img src='data:image/png;base64,"
        f"{base64.b64encode(vault.read(Path(s['path']))).decode()}'>"
        f"<figcaption>{html.escape(s['label'])} &middot; {s['ts'][11:]}</figcaption></figure>"
        for s in shots if Path(s["path"]).exists())
    return f"""<!doctype html><meta charset="utf-8"><style>
body{{font:11px/1.5 Arial,sans-serif;color:#1B2128;margin:0;padding:26px 30px}}
h1{{font-size:19px;margin:0;color:#123A6B}}
.sub{{color:#5E6B78;font-size:11px;margin:3px 0 18px}}
h2{{font-size:11px;letter-spacing:.1em;text-transform:uppercase;color:#123A6B;
border-bottom:1.5px solid #123A6B;padding-bottom:4px;margin:20px 0 9px}}
table{{width:100%;border-collapse:collapse;font-size:10.5px}}
th{{text-align:left;background:#EEF2F6;padding:5px 7px;border-bottom:1px solid #C6CFD8}}
td{{padding:4px 7px;border-bottom:1px solid #E6EBF1;vertical-align:top}}
td.n{{font-family:ui-monospace,monospace;text-align:right;white-space:nowrap}}
.s-done{{color:#146332;font-weight:700}}.s-fail{{color:#8E231B;font-weight:700}}
.s-skipped{{color:#6A7684}}
.kv{{display:flex;flex-wrap:wrap;gap:4px 28px}}
.kv div{{min-width:160px}}.kv b{{display:block;color:#5E6B78;font-weight:400;font-size:9.5px;
letter-spacing:.09em;text-transform:uppercase}}
.le{{font-family:ui-monospace,monospace;font-size:9px;border-bottom:1px solid #F0F3F6;padding:1.5px 0}}
.le span{{color:#8A96A2;margin-right:9px}}
figure{{margin:0 0 10px;break-inside:avoid}}
figure img{{width:100%;border:1px solid #C6CFD8}}
figcaption{{font-size:9px;color:#5E6B78;padding-top:3px}}
</style>
<h1>Agent run report &middot; run {run_id}</h1>
<div class="sub">Generated {datetime.now():%d %b %Y %H:%M} &middot; Solara, demo environment
&middot; sandbox portal</div>
<div class="kv">
<div><b>Customer</b>{c['name']} ({c['id']})</div>
<div><b>Consumer number</b>{c['consumer_no'] or 'not yet allotted'}</div>
<div><b>Capacity</b>{c['kw']:g} kW</div>
<div><b>Workflow</b>{r['workflow']}</div>
<div><b>Status</b>{r['status']}</div>
<div><b>Application number</b>{r['application_no'] or 'not filed'}</div>
<div><b>Started by</b>{r['started_by']}</div>
<div><b>Started</b>{(r['started_at'] or '').replace('T', ' ')}</div>
<div><b>Finished</b>{(r['finished_at'] or '-').replace('T', ' ')}</div>
<div><b>Retries</b>{r['retries']}</div>
</div>
<h2>Steps</h2><table><tr><th>#</th><th>Step</th><th>Status</th><th>Note</th><th>Time</th></tr>
{rows}</table>
<h2>Human checkpoints</h2><table><tr><th>Checkpoint</th><th>Status</th><th>Resolved by</th>
<th>Decision</th><th>At</th></tr>{hrows}</table>
<h2>Evidence</h2>{imgs}
<h2>Log</h2>{log}"""


async def run_report_pdf(run_id):
    path = EXPORTS / f"solara_run{run_id}_{datetime.now():%Y%m%d_%H%M}.pdf"
    async with async_playwright() as p:
        b = await p.chromium.launch()
        page = await b.new_page()
        await page.set_content(_run_report_html(run_id), wait_until="load")
        await page.pdf(path=str(path), format="A4", print_background=True,
                       margin={"top": "0", "bottom": "14mm", "left": "0", "right": "0"},
                       display_header_footer=True,
                       footer_template="<div style='font:8px Arial;color:#8A96A2;width:100%;"
                                       "padding:0 30px;display:flex;justify-content:space-between'>"
                                       "<span>Solara agent run report</span>"
                                       "<span class='pageNumber'></span></div>",
                       header_template="<div></div>")
        await b.close()
    return path
